#!/usr/bin/env python3
"""Import spicycenter_global_keyword_universe_100k.xlsx into data/keyword-universe/."""
from __future__ import annotations

import gzip
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Downloads" / "spicycenter_global_keyword_universe_100k.xlsx"
OUT = ROOT / "data" / "keyword-universe"

SPECIAL = {
    "United States": "usa",
    "United Kingdom": "uk",
    "United Arab Emirates": "uae",
    "Saudi Arabia": "saudi-arabia",
    "South Korea": "south-korea",
    "South Africa": "south-africa",
    "New Zealand": "new-zealand",
    "Czech Republic": "czech-republic",
    "Sri Lanka": "sri-lanka",
    "Hong Kong": "hong-kong",
}
PRODUCT_SLUG = {
    "indian spices": "indian-spices",
    "bulk spices": "bulk-spices",
    "wholesale spices": "wholesale-spices",
    "turmeric": "turmeric",
    "turmeric powder": "turmeric-powder",
    "cumin seeds": "cumin-seeds",
    "cumin powder": "cumin-powder",
}


def slugify(name: str) -> str:
    if name in SPECIAL:
        return SPECIAL[name]
    return re.sub(r"[^a-z0-9]+", "-", name.lower().strip()).strip("-")


def main() -> None:
    import openpyxl

    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not src.exists():
        raise SystemExit(f"Workbook not found: {src}")
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    markets = []
    for i, row in enumerate(wb["Country_Targets"].iter_rows(values_only=True)):
        if i == 0:
            continue
        country, code, rec = row
        markets.append(
            {
                "name": country,
                "slug": slugify(country),
                "marketCode": code,
                "recommendedContent": rec,
                "keywordCount": 0,
            }
        )

    seeds = []
    for i, row in enumerate(wb["Product_Content_Seeds"].iter_rows(values_only=True)):
        if i == 0:
            continue
        seeds.append({"name": row[0], "contentRequired": row[1], "slug": slugify(row[0])})

    gov = []
    for i, row in enumerate(wb["SEO_Governance"].iter_rows(values_only=True)):
        if i == 0:
            continue
        gov.append({"rule": row[0], "guidance": row[1]})

    by_country = defaultdict(lambda: {"count": 0, "products": defaultdict(int), "intents": defaultdict(int), "samples": []})
    by_product = defaultdict(lambda: {"count": 0, "countries": defaultdict(int), "intents": defaultdict(int), "samples": []})
    clusters = defaultdict(lambda: {"count": 0, "intents": defaultdict(int), "samples": []})

    n = 0
    raw_lines = []
    for i, row in enumerate(wb["Keyword_Master_100K"].iter_rows(values_only=True)):
        if i == 0:
            continue
        kw, prod, intent, _mod, country, _lt, _tm, _lp, _pr, _st = row
        pslug = PRODUCT_SLUG.get(prod, slugify(prod or "unknown"))
        mslug = slugify(country or "unknown")
        rec = {"k": kw, "p": pslug, "i": intent, "m": mslug, "c": country}
        raw_lines.append(json.dumps(rec, ensure_ascii=False) + "\n")
        n += 1
        bc = by_country[mslug]
        bc["count"] += 1
        bc["products"][pslug] += 1
        bc["intents"][intent] += 1
        if len(bc["samples"]) < 8:
            bc["samples"].append(kw)
        bp = by_product[pslug]
        bp["count"] += 1
        bp["name"] = prod
        bp["countries"][mslug] += 1
        bp["intents"][intent] += 1
        if len(bp["samples"]) < 8:
            bp["samples"].append(kw)
        cl = clusters[f"{pslug}::{mslug}"]
        cl["count"] += 1
        cl["intents"][intent] += 1
        if len(cl["samples"]) < 6:
            cl["samples"].append(kw)
    wb.close()

    for m in markets:
        extra = by_country.get(m["slug"], {})
        m["keywordCount"] = extra.get("count", 0)
        m["productCounts"] = dict(extra.get("products", {}))
        m["intentCounts"] = dict(extra.get("intents", {}))
        m["sampleKeywords"] = extra.get("samples", [])

    products = [
        {
            "slug": slug,
            "name": extra["name"],
            "keywordCount": extra["count"],
            "countryCounts": dict(extra["countries"]),
            "intentCounts": dict(extra["intents"]),
            "sampleKeywords": extra["samples"],
        }
        for slug, extra in sorted(by_product.items(), key=lambda x: -x[1]["count"])
    ]
    cluster_list = []
    for key, extra in clusters.items():
        p, m = key.split("::")
        cluster_list.append(
            {
                "productSlug": p,
                "marketSlug": m,
                "keywordCount": extra["count"],
                "intentCounts": dict(extra["intents"]),
                "sampleKeywords": extra["samples"],
                "url": f"/markets/{m}",
            }
        )

    (OUT / "meta.json").write_text(
        json.dumps(
            {
                "source": src.name,
                "importedAt": "2026-09-19",
                "keywordCount": n,
                "validation": "Generated seed — not claimed as search volume",
                "governance": gov,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    (OUT / "markets.json").write_text(json.dumps(markets, indent=2), encoding="utf-8")
    (OUT / "products.json").write_text(json.dumps(products, indent=2), encoding="utf-8")
    (OUT / "clusters.json").write_text(json.dumps(cluster_list, indent=2), encoding="utf-8")
    (OUT / "seeds.json").write_text(json.dumps(seeds, indent=2), encoding="utf-8")
    with gzip.open(OUT / "keywords.jsonl.gz", "wt", encoding="utf-8") as fh:
        fh.writelines(raw_lines)
    print(f"Wrote {n} keywords to {OUT}")


if __name__ == "__main__":
    main()
